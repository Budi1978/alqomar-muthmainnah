import re, os, json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HIJAU=RGBColor(0x1B,0x4B,0x2E)
def shade(cell, hexcolor):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd')
    shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),hexcolor); tcPr.append(shd)
def add_runs(par, text, size=10):
    # **bold** inline
    parts=re.split(r'(\*\*[^*]+\*\*)', text)
    for p in parts:
        if not p: continue
        if p.startswith('**') and p.endswith('**'):
            r=par.add_run(p[2:-2]); r.bold=True
        else:
            r=par.add_run(p.replace('~~',''))
        r.font.size=Pt(size)
def base(doc):
    st=doc.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(10.5)
    for s in doc.sections:
        s.left_margin=s.right_margin=Cm(1.8); s.top_margin=s.bottom_margin=Cm(1.6)
def heading(doc, text, level=1):
    h=doc.add_heading(text.replace('**',''), level=level)
    for r in h.runs: r.font.color.rgb=HIJAU
    return h
def table(doc, rows, widths=None, size=9):
    if not rows: return
    ncol=max(len(r) for r in rows)
    t=doc.add_table(rows=0, cols=ncol); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        for j in range(ncol):
            txt=row[j] if j<len(row) else ''
            c=cells[j]; c.text=''
            p=c.paragraphs[0]; add_runs(p, txt, size)
            if i==0:
                for r in p.runs: r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
                shade(c,'1B4B2E')
            elif j==0: 
                for r in p.runs: r.bold=True
                shade(c,'F2F6F2')
    if widths:
        for row in t.rows:
            for j,w in enumerate(widths):
                if j<len(row.cells): row.cells[j].width=Cm(w)
    doc.add_paragraph()
    return t
def md_table_rows(lines):
    rows=[]
    for ln in lines:
        if re.match(r'^\|[\s:|-]+\|$', ln.strip()): continue
        cells=[c.strip() for c in ln.strip().strip('|').split('|')]
        rows.append(cells)
    return rows
def md_to_docx(md, doc, widths_by_cols=None):
    lines=md.split('\n'); i=0
    while i<len(lines):
        ln=lines[i]
        if ln.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].startswith('|'): block.append(lines[i]); i+=1
            rows=md_table_rows(block); n=len(rows[0]) if rows else 0
            table(doc, rows, (widths_by_cols or {}).get(n))
            continue
        if ln.startswith('# '): heading(doc, ln[2:], 0 if False else 1)
        elif ln.startswith('## '): heading(doc, ln[3:], 2)
        elif ln.startswith('### '): heading(doc, ln[4:], 3)
        elif re.match(r'^\s*[-•] ', ln): p=doc.add_paragraph(style='List Bullet'); add_runs(p, re.sub(r'^\s*[-•] ','',ln))
        elif re.match(r'^\s*\d+\. ', ln): p=doc.add_paragraph(style='List Number'); add_runs(p, re.sub(r'^\s*\d+\. ','',ln))
        elif ln.strip()=='---': pass
        elif ln.strip(): p=doc.add_paragraph(); add_runs(p, ln)
        i+=1

# ---------- 1. MASTER ----------

# ---------- konversi panduan SMP ----------
import sys, glob
src_dir, dst_dir = sys.argv[1], sys.argv[2]
os.makedirs(dst_dir, exist_ok=True)
W={2:[5.5,11.5],3:[4.0,7.5,5.5],4:[2.0,4.0,7.5,3.5],5:[1.8,3.5,3.5,3.5,4.7],6:[1.2,2.6,6.0,2.2,1.2,3.8]}
for md_path in sorted(glob.glob(os.path.join(src_dir,'PANDUAN-*.md'))):
    md=open(md_path).read()
    doc=Document(); base(doc)
    md_to_docx(md, doc, W)
    name=os.path.basename(md_path)[len('PANDUAN-'):-3]
    out=os.path.join(dst_dir, f'Panduan-Menurunkan-CP-menjadi-ATP-{name}-SMP.docx')
    doc.save(out)
    words=len(re.sub(r'[|#*-]',' ',md).split()); tables=len(re.findall(r'\n\|[^\n]*\n\|[\s:|-]+\|', md))
    print(f'{os.path.basename(out)}: {words} kata, {tables} tabel')
