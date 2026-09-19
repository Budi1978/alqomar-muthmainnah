import re, os, json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HIJAU=RGBColor(0x1B,0x4B,0x2E)
def shade(cell, hexcolor):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd')
    shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),hexcolor); tcPr.append(shd)
def add_runs(par, text, size=10):
    # **bold** inline
    parts=re.split(r'(\*\*[^*]+\*\*)', text)
    for p in parts:
        if not p: continue
        if p.startswith('**') and p.endswith('**'):
            r=par.add_run(p[2:-2]); r.bold=True
        else:
            r=par.add_run(p.replace('~~',''))
        r.font.size=Pt(size)
def base(doc):
    st=doc.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(10.5)
    for s in doc.sections:
        s.left_margin=s.right_margin=Cm(1.8); s.top_margin=s.bottom_margin=Cm(1.6)
def heading(doc, text, level=1):
    h=doc.add_heading(text.replace('**',''), level=level)
    for r in h.runs: r.font.color.rgb=HIJAU
    return h
def table(doc, rows, widths=None, size=9):
    if not rows: return
    ncol=max(len(r) for r in rows)
    t=doc.add_table(rows=0, cols=ncol); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        for j in range(ncol):
            txt=row[j] if j<len(row) else ''
            c=cells[j]; c.text=''
            p=c.paragraphs[0]; add_runs(p, txt, size)
            if i==0:
                for r in p.runs: r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
                shade(c,'1B4B2E')
            elif j==0: 
                for r in p.runs: r.bold=True
                shade(c,'F2F6F2')
    if widths:
        for row in t.rows:
            for j,w in enumerate(widths):
                if j<len(row.cells): row.cells[j].width=Cm(w)
    doc.add_paragraph()
    return t
def md_table_rows(lines):
    rows=[]
    for ln in lines:
        if re.match(r'^\|[\s:|-]+\|$', ln.strip()): continue
        cells=[c.strip() for c in ln.strip().strip('|').split('|')]
        rows.append(cells)
    return rows
def md_to_docx(md, doc, widths_by_cols=None):
    lines=md.split('\n'); i=0
    while i<len(lines):
        ln=lines[i]
        if ln.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].startswith('|'): block.append(lines[i]); i+=1
            rows=md_table_rows(block); n=len(rows[0]) if rows else 0
            table(doc, rows, (widths_by_cols or {}).get(n))
            continue
        if ln.startswith('# '): heading(doc, ln[2:], 0 if False else 1)
        elif ln.startswith('## '): heading(doc, ln[3:], 2)
        elif ln.startswith('### '): heading(doc, ln[4:], 3)
        elif re.match(r'^\s*[-•] ', ln): p=doc.add_paragraph(style='List Bullet'); add_runs(p, re.sub(r'^\s*[-•] ','',ln))
        elif re.match(r'^\s*\d+\. ', ln): p=doc.add_paragraph(style='List Number'); add_runs(p, re.sub(r'^\s*\d+\. ','',ln))
        elif ln.strip()=='---': pass
        elif ln.strip(): p=doc.add_paragraph(); add_runs(p, ln)
        i+=1

# ---------- 1. MASTER ----------
master=open('tugas-kurikulum/REKAP-Pemeriksaan-ATP-24-guru-2026-09-19.md').read()
doc=Document(); base(doc)
md_to_docx(master, doc, {4:[3.2,5.0,5.0,4.2],3:[3.0,5.5,9.0],2:[7,10]})
doc.save('tugas-kurikulum/word/REKAP-Pemeriksaan-ATP-24-Guru-2026-09-19.docx')
print('master ok')

# ---------- 2. PER GURU ----------
def parse_sections(path):
    txt=open(path).read(); secs=[]; cur=None
    for ln in txt.split('\n'):
        if ln.startswith('## '):
            if cur: secs.append(cur)
            cur={'judul':ln[3:].strip(),'verdict':'','ringkas':'','rows':[],'lain':[]}
        elif cur is not None:
            if ln.startswith('**Verdict'): cur['verdict']=re.sub(r'\*\*Verdict:?\*\*:?','',ln).replace('**','').strip(' :')
            elif ln.startswith('**Ringkas'): cur['ringkas']=re.sub(r'\*\*Ringkas[^*]*\*\*:?','',ln).strip()
            elif ln.startswith('|'): cur['rows'].append(ln)
            elif ln.strip(): cur['lain'].append(ln)
    if cur: secs.append(cur)
    return [s for s in secs if s['rows'] and not s['judul'].lower().startswith(('rekap','tabel rekap','kesimpulan'))]
S={}
for g in ['review-G1.md','review-G2.md','review-G3.md','review-G4.md']: S[g]=parse_sections(g)
S['ulfa']=parse_sections('tugas-kurikulum/koreksi-ATP-Ulfa-2026-09-19.md')
def find(g, key): 
    r=[s for s in S[g] if key.lower() in s['judul'].lower()]
    return r
# teacher -> list of (section, catatan kelompok)
T={}
def add(nama, secs, note=''):
    T.setdefault(nama, {'secs':[], 'note':note})
    for s in secs: T[nama]['secs'].append(s)
    if note: T[nama]['note']=note
add("Een Muflihat, S.Pd", find('review-G1.md','Een Muflihat'), "ATP kelas 1 disusun bersama Usth. Sa'adiyah Ulfa (wali kelas paralel). Koreksi ini berlaku untuk keduanya.")
add("Sa'adiyah Ulfa, S.Pd", S['ulfa'], "ATP kelas 1 disusun bersama Usth. Een Muflihat (wali kelas paralel). Koreksi ini berlaku untuk keduanya.")
add("Kustiah, S.Pd", find('review-G1.md','Kustiah'), "ATP Bahasa Indonesia kelas 2 disusun bersama Usth. Ningsih Nurna; koreksinya berlaku untuk keduanya.")
add("Ningsih Nurna, S.Pd", find('review-G1.md','Ningsih')+find('review-G1.md','Kustiah, S.Pd — Bahasa Indonesia'), "ATP Bahasa Indonesia kelas 2 disusun bersama Usth. Kustiah; bagian kedua di bawah adalah koreksi dokumen induknya.")
add("Rifatul Hasanah, S.Pd", find('review-G1.md','Rifatul'))
add("Iman Paojan, S.Pd.I", find('review-G1.md','Iman Paojan'), "Tugas hanya mencakup semester ganjil; catatan tentang semester genap di tabel boleh diabaikan.")
add("Syafsilaroza Octaria, S.Pd", find('review-G2.md','Syafsilaroza'), "ATP kelas 3 berlaku juga untuk Dra. Sjarniwati (wali kelas paralel), yang belum mengumpulkan.")
fc=find('review-G2.md','Kelas VI / Fase C')+find('review-G2.md','Pendidikan Pancasila, Kelas V')
for n in ["Ellida Siregar, S.Pd","Selvi Nur Fitriah, S.Hum","Fatmarianti, S.Pd"]:
    add(n, fc, "ATP Fase C disusun bersama (Ellida, Selvi, Fatmarianti). Koreksi berlaku untuk ketiganya. Catatan khusus Fatmarianti: kelas 6 perlu ATP tersendiri (buku dan bab kelas 6 berbeda dari kelas 5).")
add("Yushal Rachmat Gumilar, S.Pd", find('review-G2.md','Yushal'))
add("Rahayu Vina Purwanti, S.Pd", find('review-G2.md','Rahayu'))
add("Widya, S.Pd", find('review-G2.md','Widya'))
add("Amanda Raihan, S.Pd", find('review-G3.md','Amanda'))
add("Bidadari Sholihah, S.Pd", find('review-G3.md','Bidadari'))
add("Husnah, S.Pd", find('review-G3.md','Husnah'))
add("Iis Apriyanti, S.Pd", find('review-G3.md','Iis'))
add("Hernita Umardi, S.Pd", find('review-G3.md','Hernita'))
add("Sulistiah Sari, S.Pd", find('review-G3.md','Sulistiah'))
add("Ropiyati, S.Pd.I", find('review-G4.md','Ropiyati'))
add("Rosdiana, S.Pd.AUD", find('review-G4.md','Rosdiana'))
add("Mita Yulianah, S.Pd", find('review-G4.md','Mita'))
add("Dimiyati Dyas Anindita, S.Pd", find('review-G4.md','Dimiyati'), "Tugas hanya mencakup semester ganjil; catatan tentang semester genap boleh diabaikan.")
add("Siti Salamatul Laila, S.Pd", find('review-G4.md','Siti Salamatul'))

# CP rows from 4c
m=re.search(r'## 4c\..*?\n(\|.*?)\n\n\*\*Kesimpulan', master, re.S)
cprows=md_table_rows(m.group(1).split('\n'))[1:]
CPKEY={"Een Muflihat, S.Pd":['kls 1 (Een'],"Sa'adiyah Ulfa, S.Pd":['kls 1 (Een'],"Kustiah, S.Pd":['Kustiah'],"Ningsih Nurna, S.Pd":['Kustiah–Ningsih'],
 "Rifatul Hasanah, S.Pd":['Rifatul'],"Iman Paojan, S.Pd.I":['Iman'],"Syafsilaroza Octaria, S.Pd":['Syafsilaroza'],
 "Ellida Siregar, S.Pd":['Selvi–Fatma','Ellida–Selvi'],"Selvi Nur Fitriah, S.Hum":['Selvi–Fatma','Ellida–Selvi'],"Fatmarianti, S.Pd":['Selvi–Fatma','Ellida–Selvi'],
 "Yushal Rachmat Gumilar, S.Pd":['Yushal'],"Rahayu Vina Purwanti, S.Pd":['Rahayu'],"Widya, S.Pd":['Widya'],"Amanda Raihan, S.Pd":['Amanda'],
 "Bidadari Sholihah, S.Pd":['Bidadari'],"Husnah, S.Pd":['Husnah'],"Iis Apriyanti, S.Pd":['Iis'],"Hernita Umardi, S.Pd":['Hernita'],"Sulistiah Sari, S.Pd":['Sulistiah'],
 "Ropiyati, S.Pd.I":['Ropiyati'],"Rosdiana, S.Pd.AUD":['Rosdiana'],"Mita Yulianah, S.Pd":['Mita'],"Dimiyati Dyas Anindita, S.Pd":['Dimiyati'],"Siti Salamatul Laila, S.Pd":['Siti Laila']}


# ===== LAMPIRAN per guru: CP resmi + panduan contoh =====
CPJ=json.load(open('cp-2025-parsed.json'))
_p=re.sub(r'\s+',' ',open('cp-mapel/PAUD.txt').read().replace('\xa0',' '))
_i=_p.rfind('Capaian Pembelajaran Fase Fondasi yang terdiri dari tiga elemen'); _j=_p.find('KEPALA BADAN',_i)
PAUD_TXT=_p[_i:_j].replace('​','')
GURU_MAPEL={
 "Een Muflihat, S.Pd":[('PANCASILA','A','Pendidikan Pancasila Fase A (kelas 1–2)'),('BINDO','A','Bahasa Indonesia Fase A (kelas 1–2)')],
 "Sa'adiyah Ulfa, S.Pd":[('PANCASILA','A','Pendidikan Pancasila Fase A (kelas 1–2)'),('BINDO','A','Bahasa Indonesia Fase A (kelas 1–2)')],
 "Kustiah, S.Pd":[('BINDO','A','Bahasa Indonesia Fase A (kelas 1–2)'),('PANCASILA','A','Pendidikan Pancasila Fase A (kelas 1–2)')],
 "Ningsih Nurna, S.Pd":[('BINDO','A','Bahasa Indonesia Fase A (kelas 1–2)')],
 "Rifatul Hasanah, S.Pd":[('PAI','A','PAI dan Budi Pekerti Fase A (kelas 1–2)'),('PAI','C','PAI dan Budi Pekerti Fase C (kelas 5–6)')],
 "Iman Paojan, S.Pd.I":[('PAI','A','PAI dan Budi Pekerti Fase A (kelas 1–2)'),('PAI','C','PAI dan Budi Pekerti Fase C (kelas 5–6)')],
 "Syafsilaroza Octaria, S.Pd":[('BINDO','B','Bahasa Indonesia Fase B (kelas 3–4)')],
 "Ellida Siregar, S.Pd":[('BINDO','C','Bahasa Indonesia Fase C (kelas 5–6)'),('PANCASILA','C','Pendidikan Pancasila Fase C (kelas 5–6)')],
 "Selvi Nur Fitriah, S.Hum":[('BINDO','C','Bahasa Indonesia Fase C (kelas 5–6)'),('PANCASILA','C','Pendidikan Pancasila Fase C (kelas 5–6)')],
 "Fatmarianti, S.Pd":[('BINDO','C','Bahasa Indonesia Fase C (kelas 5–6)'),('PANCASILA','C','Pendidikan Pancasila Fase C (kelas 5–6)')],
 "Yushal Rachmat Gumilar, S.Pd":[('BING','C','Bahasa Inggris Fase C (kelas 5–6)')],
 "Rahayu Vina Purwanti, S.Pd":[('BARAB','-','Bahasa Arab')],
 "Widya, S.Pd":[('BING','B','Bahasa Inggris Fase B (kelas 3–4)'),('BING','D','Bahasa Inggris Fase D (SMP)')],
 "Amanda Raihan, S.Pd":[('BINDO','D','Bahasa Indonesia Fase D (SMP)')],
 "Bidadari Sholihah, S.Pd":[('IPA','D','IPA Fase D (SMP)'),('IPS','D','IPS Fase D (SMP)')],
 "Husnah, S.Pd":[('MAT','D','Matematika Fase D (SMP)')],
 "Iis Apriyanti, S.Pd":[('IPA','D','IPA Fase D (SMP)')],
 "Hernita Umardi, S.Pd":[('PAI','D','PAI dan Budi Pekerti Fase D (SMP)')],
 "Sulistiah Sari, S.Pd":[('TAHFIDZ','-','Tahfidz')],
 "Ropiyati, S.Pd.I":[('PAUD','Fondasi','PAUD Fase Fondasi')],
 "Rosdiana, S.Pd.AUD":[('PAUD','Fondasi','PAUD Fase Fondasi')],
 "Mita Yulianah, S.Pd":[('PAUD','Fondasi','PAUD Fase Fondasi')],
 "Dimiyati Dyas Anindita, S.Pd":[('PAUD','Fondasi','PAUD Fase Fondasi'),('BINDO','D','Bahasa Indonesia Fase D (SMP)')],
 "Siti Salamatul Laila, S.Pd":[('INFORMATIKA','D','Informatika Fase D (SMP)'),('INFORMATIKA','SD','Informatika SD')],
}
TANPA_CP={'BARAB':"Bahasa Arab tidak memiliki CP di Kepka BSKAP 046/H/KR/2025 untuk SD/SMP (yang ada hanya Fase F untuk SMA). Tuliskan di Tahap 1: \"Sumber: <dokumen yang dipakai sekolah, mis. KMA Kemenag tentang kurikulum Bahasa Arab MI/MTs, atau target program sekolah>\", lalu salin capaian per elemen dari dokumen itu. Kepala Sekolah/koordinator akan menetapkan dokumen mana yang dipakai.",
 'TAHFIDZ':"Tahfidz tidak memiliki CP nasional. Tuliskan di Tahap 1: \"Sumber: Target Program Tahfidz Al-Qomar Muthmainnah TA 2026/2027\" dan salin target hafalan per kelas (surah, jumlah ayat/juz) sebagai pengganti bunyi CP. Elemen dapat memakai: Ziyadah (hafalan baru), Muroja'ah, Tahsin/Tajwid, Adab.",
 'SD':"Informatika di Kepka BSKAP 046/H/KR/2025 hanya tersedia untuk Fase D–F (SMP–SMA); tidak ada CP Informatika untuk SD. Untuk kelas 1, tuliskan di Tahap 1: \"Sumber: dokumen target sekolah (muatan lokal/ekstra)\" dan capaian yang ditetapkan sekolah. Kepala Sekolah akan menetapkan dokumen rujukannya."}
def lampiran_cp(doc, nama):
    heading(doc,'Lampiran A — Bunyi CP resmi (Kepka BSKAP 046/H/KR/2025) untuk mapel dan fase Anda',1)
    doc.add_paragraph('Salin teks di bawah ini apa adanya ke Tahap 1, satu baris per elemen, lalu tulis "Sumber: Kepka BSKAP No. 046/H/KR/2025". Jangan mengubah kata-katanya.')
    for mapel,fase,label in GURU_MAPEL.get(nama,[]):
        heading(doc,label,2)
        if fase in TANPA_CP and mapel in ('BARAB','TAHFIDZ') or fase=='SD':
            doc.add_paragraph(TANPA_CP.get(mapel, TANPA_CP.get(fase)))
            continue
        if mapel=='PAUD':
            doc.add_paragraph('CP PAUD Fase Fondasi terdiri dari 3 elemen; tiap elemen punya beberapa subelemen (tanda ●). Salin SEMUA subelemen dari elemen yang menjadi fokus ATP Anda, bukan hanya satu kalimat.')
            for seg in re.split(r'\s(?=\d\.\s(?:Nilai Agama|Jati Diri|Dasar-dasar))', PAUD_TXT):
                seg=seg.strip()
                if not seg or seg.startswith('Capaian Pembelajaran Fase Fondasi yang'): continue
                parts=re.split(r'\s●\s*', seg)
                p=doc.add_paragraph(); add_runs(p, '**'+parts[0].strip()+'**', 10)
                for b in parts[1:]:
                    p=doc.add_paragraph(style='List Bullet'); add_runs(p, b.strip(), 10)
            continue
        els=CPJ.get(mapel,{}).get(fase,[])
        if not els: doc.add_paragraph('(teks CP tidak tersedia dalam berkas ini — salin dari PDF resmi)'); continue
        rows=[['Elemen','Bunyi CP akhir fase (salin persis)']]+[[n,t.replace('​','')] for n,t in els]
        table(doc, rows, [4.5,12.5], 9.5)
PANDUAN=[
 ("Memecah TP yang memuat dua kemampuan",
  "SALAH: \"Peserta didik dapat mengklasifikasikan makhluk hidup serta mengevaluasi dasar pengelompokannya.\"\nBENAR: TP 2a \"Peserta didik dapat mengklasifikasikan makhluk hidup berdasarkan ciri yang diamati.\"  TP 2b \"Peserta didik dapat menilai ketepatan dasar pengelompokan makhluk hidup.\"\nAturan: satu TP = satu kata kerja utama = satu hal yang bisa dinilai."),
 ("Mengganti kata kerja yang tidak terukur",
  "Hindari: memahami, mengetahui, mengenal, menyadari. Ganti dengan yang bisa dilihat hasilnya: menyebutkan, menunjukkan, menjelaskan, membandingkan, menghitung, mengurutkan, menceritakan kembali, mempraktikkan, membuat.\nSALAH: \"Peserta didik memahami aturan di rumah.\"  BENAR: \"Peserta didik dapat menyebutkan tiga aturan di rumah dan menceritakan akibat jika dilanggar.\""),
 ("Memisahkan 'dengan cara …' dari TP",
  "TP berisi tujuan, bukan kegiatan. Bagian \"dengan cara mengamati gambar dan bermain peran\" dipindahkan ke modul ajar/RPP.\nSEBELUM: \"Peserta didik dapat menyebutkan simbol Garuda Pancasila dengan cara mengamati gambar dan menunjuk tiap simbolnya.\"  SESUDAH: \"Peserta didik dapat menyebutkan simbol-simbol pada lambang negara Garuda Pancasila.\""),
 ("Menandai bagian CP yang tidak ada di buku",
  "Baca CP kalimat per kalimat. Setiap kalimat harus punya minimal satu TP. Kalau bukunya tidak memuat materi itu, tetap tulis TP-nya, isi kolom Bab dengan \"—\" dan kolom Catatan dengan \"PERLU MATERI TAMBAHAN: <sumber yang akan dipakai>\". Jangan dibiarkan kosong, dan jangan menjawab \"Ya\" di Tahap 6 nomor 1 kalau masih ada yang seperti ini."),
 ("Mengisi JP dan kolom Semester dengan checkpoint",
  "Hitung dulu jam nyata: JP per minggu × jumlah minggu efektif (semester ganjil ± 17 minggu, dikurangi pekan asesmen). Contoh 4 JP/minggu → ± 60 JP. Bagi ke TP sesuai berat materinya (TP praktik ringan 2–3 JP, TP konsep inti 6–9 JP). Kolom Semester diisi dengan checkpoint kalender, contoh: TP 1–3 \"Ganjil – sblm ASB1\", TP 4–6 \"Ganjil – sblm ASTS\", TP 7–9 \"Ganjil – sblm ASB2\", TP 10–12 \"Ganjil – sblm ASAS\". Tambahkan baris terakhir: TOTAL JP GANJIL = … dan bandingkan dengan jam nyata; selisih besar berarti ada TP yang kurang atau JP yang terlalu kecil."),
 ("Mengurutkan TP sesuai prasyarat",
  "Bab yang di Tahap 3 dijawab \"Tidak ada\" (gerbang) diletakkan paling awal. TP yang butuh TP lain ditulis nomor prasyaratnya di kolom Prasyarat, misalnya \"TP 1, TP 2\", bukan \"Bab 1\". Bab yang bebas dipakai mengisi celah dekat pekan asesmen, bukan bab berat."),
 ("Menjawab Tahap 6 dengan jujur",
  "\"Belum pasti\" adalah jawaban yang benar bila memang ada bagian CP tanpa materi, atau ATP kelas berikutnya belum ada. Tulis alasannya konkret: \"Belum pasti: bagian CP tentang aturan di sekolah belum ada di buku (TP 8, 16), akan memakai modul internal.\" Jawaban \"Ya\" tanpa bukti akan dikembalikan."),
 ("Menulis identitas dan sumber",
  "Nama guru, mapel, kelas/fase harus sama di semua kepala halaman (Tahap 1 sampai 6). ATP yang disusun bersama wali kelas paralel: tulis kedua nama di setiap kepala halaman dan keduanya menandatangani. Di bawah tabel Tahap 1 tulis: \"Sumber: Kepka BSKAP No. 046/H/KR/2025, CP <mapel> Fase <X>\"."),
]
def lampiran_panduan(doc):
    heading(doc,'Lampiran B — Panduan cara memperbaiki (contoh sebelum–sesudah)',1)
    doc.add_paragraph('Gunakan contoh berikut saat mengerjakan poin revisi di bagian 2 dan 4. Contoh diambil dari kesalahan yang paling sering muncul di seluruh tugas.')
    for judul,isi in PANDUAN:
        heading(doc, judul, 2)
        for ln in isi.split('\n'):
            p=doc.add_paragraph(); add_runs(p, ln, 10)

def per_guru(nama, info):
    doc=Document(); base(doc)
    h=doc.add_heading('Koreksi Lembar Kerja Penyusunan ATP', 0)
    for r in h.runs: r.font.color.rgb=HIJAU; r.font.size=Pt(20)
    p=doc.add_paragraph(); add_runs(p, f"**Nama guru:** {nama}", 11)
    p=doc.add_paragraph(); add_runs(p, "**Tugas:** Pelatihan Guru — Lembar Kerja Penyusunan ATP (semester ganjil 2026/2027), dikumpulkan 19 September 2026", 10.5)
    p=doc.add_paragraph(); add_runs(p, "**Pemeriksa:** Kepala Sekolah SDIT/SMPIT/TKIT Al-Qomar Muthmainnah", 10.5)
    if info['note']:
        p=doc.add_paragraph(); add_runs(p, "**Catatan:** "+info['note'], 10)
    heading(doc,'1. Ringkasan hasil',1)
    rows=[['Dokumen','Verdict','Inti masalah']]
    for s in info['secs']:
        jud=re.sub(r'\s*\([^)]*\.txt[^)]*\)','',s['judul']); jud=re.sub(r'^[A-Z]\.\s','',jud)
        v=re.sub(r'^Verdict:?\s*','',s['verdict']); vv=v.split('.',1); rows.append([jud, vv[0].strip(' .'), s['ringkas'] or (vv[1].strip() if len(vv)>1 and vv[1].strip() else '—')])
    table(doc, rows, [6,3.5,7.5])
    p=doc.add_paragraph(); add_runs(p, "Arti verdict: **LAYAK, REVISI KECIL** = dokumen diterima, perbaiki poin yang disebut lalu kumpulkan ulang. **PERLU REVISI** = ada kesalahan mendasar (CP, identitas, atau ATP final), dokumen dikembalikan untuk diperbaiki sebelum ditandatangani.", 9.5)
    heading(doc,'2. Kesalahan per tahap dan cara menyelesaikannya',1)
    p=doc.add_paragraph(); add_runs(p,'Tabel mengikuti 6 tahap lembar kerja. Jika ada kalimat "cocokkan ke 046/2025", hasil pastinya ada di bagian 3 (sudah dicocokkan dengan dokumen resmi).',9.5)
    langkah=[]
    for s in info['secs']:
        jud=re.sub(r'\s*\([^)]*\.txt[^)]*\)','',s['judul']); jud=re.sub(r'^[A-Z]\.\s','',jud)
        heading(doc, jud, 2)
        rws=md_table_rows(s['rows'])
        # columns: Tahap | Kondisi | Salah/kurang | Perbaikan -> keep Tahap | Kesalahan | Cara menyelesaikan
        out=[['Tahap','Kesalahan / kekurangan','Cara menyelesaikan']]
        for r in rws[1:]:
            if len(r)>=4:
                out.append([r[0], r[2], r[3]])
                if not r[2].strip().lower().startswith('tidak ada') and r[3].strip() not in ('—','-',''): langkah.append((jud, r[0], r[3]))
            elif len(r)==3: out.append(r)
        table(doc, out, [2.6,7.0,7.4])
    heading(doc,'3. Kesesuaian CP dengan Kepka BSKAP 046/H/KR/2025',1)
    keys=CPKEY.get(nama,[]); crow=[r for r in cprows if any(k in r[0] for k in keys)]
    if crow: table(doc, [['Dokumen','Hasil pencocokan','Tindakan']]+crow, [4.5,6.5,6])
    else: doc.add_paragraph('Tidak ada catatan khusus.')
    heading(doc,'4. Urutan pengerjaan revisi',1)
    doc.add_paragraph('Kerjakan berurutan; Tahap 1 (CP) harus benar dulu karena Tahap 4–6 bergantung padanya.')
    for jud,tahap,fix in langkah:
        p=doc.add_paragraph(style='List Number'); lab=re.sub(r'\s*\(\d+ halaman\)','',jud.split(' — ')[-1] if ' — ' in jud else jud); add_runs(p, f"**{tahap}** ({lab}): {fix}", 10)
    lampiran_cp(doc, nama)
    lampiran_panduan(doc)
    doc.add_paragraph()
    p=doc.add_paragraph(); add_runs(p,"Batas pengumpulan revisi: ____________________   Diterima kembali tanggal: ____________________",10)
    doc.add_paragraph()
    t=doc.add_table(rows=2, cols=2); 
    t.cell(0,0).text='Guru yang bersangkutan'; t.cell(0,1).text='Kepala Sekolah'
    t.cell(1,0).text='\n\n\n'+nama; t.cell(1,1).text='\n\n\nUst. Ahmad Budi Setiawan, S.E.'
    for row in t.rows:
        for c in row.cells:
            for pp in c.paragraphs: pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    safe=re.sub(r'[^\w]+','-',nama).strip('-')
    fn=f'tugas-kurikulum/word/Koreksi-ATP-{safe}.docx'; doc.save(fn); return fn, len(info['secs'])
made=[]
for n,info in T.items(): made.append(per_guru(n,info))
for fn,k in made: print(k, fn)
print(len(made),'file per guru')
